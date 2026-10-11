#!/usr/bin/env python3
"""
build-conc-100-pptx.py — render deck 100: "Concurrency & Parallelism: The
Mental Model."

The general/introductory deck for engineers new to concurrency and tech
leads scoping a project — diagram-forward, light on code. Reuses every
slide builder, the theme, and the helpers from build-pptx.py — loaded via
importlib because that file's name has a hyphen — and supplies its own
title / agenda / closing slides plus this deck's SECTIONS from
sections_conc_100.py.

Output: presentation/cpp-concurrency-100-mental-model.pptx
"""
import importlib.util
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
TOOLS = ROOT / "tools"


def load_build_pptx():
    """Load build-pptx.py (hyphenated filename) as a module."""
    spec = importlib.util.spec_from_file_location(
        "build_pptx", TOOLS / "build-pptx.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


bp = load_build_pptx()
C, F, FontFam = bp.C, bp.F, bp.FontFam
SLIDE_W, SLIDE_H = bp.SLIDE_W, bp.SLIDE_H

bp.FOOTER_LABEL = "C++ Concurrency 100"

REPO_URL = "github.com/patterncatalyst/cpp-concurrency-and-parallelism"
OUT_PATH = ROOT / "presentation" / "cpp-concurrency-100-mental-model.pptx"
NOTES_PATH = ROOT / "presentation" / "build-notes-conc.md"


def build_title(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.BG_TITLE_DARK)

    dot_y = Inches(1.5)
    for i, color in enumerate([C.ACCENT_CYAN, C.ACCENT_GREEN, C.ACCENT_BLUE]):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(1.0 + i * 0.35), dot_y, Inches(0.22), Inches(0.22))
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()

    bp.add_text(slide, Inches(1.0), Inches(2.0), Inches(11.3), Inches(0.5),
                "DECK 100 — THE MENTAL MODEL", size=Pt(20), color=C.ACCENT_CYAN,
                bold=True, font=FontFam.HEADER)
    bp.add_text(slide, Inches(1.0), Inches(2.5), Inches(11.3), Inches(1.5),
                "Concurrency & Parallelism:\nThe Mental Model",
                size=F.TITLE_HUGE, color=C.TEXT_WHITE, bold=True, font=FontFam.HEADER)
    bp.add_text(slide, Inches(1.0), Inches(4.5), Inches(11.3), Inches(0.5),
                "Structure vs execution  •  the machine underneath  •  shared state  "
                "•  measuring what's true  •  concurrency under an orchestrator",
                size=F.BODY, color=C.TEXT_LIGHT, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(5.2), Inches(11.3), Inches(0.4),
                "An introduction for engineers new to concurrency, and for tech leads scoping a system",
                size=F.BODY_SMALL, color=C.TEXT_MUTED, italic=True, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(5.9), Inches(11.3), Inches(0.4),
                "Light on code, heavy on the model — the 101 and 201 decks build and tune the rest",
                size=F.BODY_SMALL, color=C.ACCENT_CYAN, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.4),
                REPO_URL, size=F.CAPTION, color=C.TEXT_MUTED, font=FontFam.CODE)

    bp.set_notes(slide,
        "This is the mental-model deck, the first of three for this book. Its "
        "job is not to teach C++ syntax; it is to install the vocabulary and "
        "the hardware intuition that every later, more code-heavy deck "
        "assumes an audience already has. Concurrency and parallelism as "
        "independent axes, the real cost structure of the machine underneath "
        "a thread, why shared state is the place correctness gets hard, how "
        "to tell a trustworthy performance number from a flattering one, and "
        "what changes once the whole program runs inside a container.\n\n"
        "The audience is engineers new to concurrency and tech leads scoping "
        "a system, so the deck stays diagram-forward and light on code on "
        "purpose. Anyone who wants the concrete standard-library mechanisms "
        "that implement these ideas moves on to the 101 deck; anyone pushing "
        "toward microsecond-scale, low-latency systems moves on to the 201 "
        "deck after that.")


def build_agenda(prs, total, sections):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.TEXT_WHITE)
    bp.add_header_bar(slide, "Overview", 2, total, "Agenda — ten ideas")
    bp.add_footer(slide, "Agenda", 2, total)

    palette = [C.ACCENT_RED, C.ACCENT_ORANGE, C.ACCENT_BLUE, C.ACCENT_GREEN,
               C.ACCENT_CYAN, C.ACCENT_PURPLE]
    rows_per_col = 5
    col_width = Inches(6.0)
    row_height = Inches(0.95)
    grid_left = Inches(0.5)
    grid_top = Inches(1.1)

    for i, sec in enumerate(sections):
        color = palette[i % len(palette)]
        col = i // rows_per_col
        row = i % rows_per_col
        x = grid_left + col * (col_width + Inches(0.3))
        y = grid_top + row * row_height

        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, Inches(0.5), Inches(0.5))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
        bp.add_text(slide, x, y, Inches(0.5), Inches(0.5), str(sec["num"]),
                    size=Pt(16), color=C.TEXT_WHITE, bold=True,
                    align=PP_ALIGN.CENTER, anchor=bp.MSO_ANCHOR.MIDDLE,
                    font=FontFam.HEADER)
        bp.add_text(slide, x + Inches(0.65), y, Inches(5.2), Inches(0.5),
                    sec["title"], size=Pt(16), color=C.TEXT_DARK, bold=True,
                    anchor=bp.MSO_ANCHOR.MIDDLE, font=FontFam.HEADER)
        bp.add_text(slide, x + Inches(0.65), y + Inches(0.42), Inches(5.2), Inches(0.5),
                    sec["tagline"], size=Pt(11), color=C.TEXT_MUTED, italic=True,
                    anchor=bp.MSO_ANCHOR.TOP, font=FontFam.BODY)

    bp.set_notes(slide,
        "Ten sections, read left to right, top to bottom. The first three "
        "build the vocabulary and the hardware ground truth: concurrency "
        "versus parallelism, the machine underneath a thread, and the clock "
        "and deadline discipline every timeout in the book depends on. The "
        "middle four get into why shared state is hard and how to reason "
        "about it: locking and deadlock, the memory model, "
        "benchmarking without deception, and the two families of "
        "inter-process communication. The last three bring it to where it "
        "runs: debugging a race or a deadlock, what changes inside "
        "a container, and where the two companion decks pick up from here.\n\n"
        "Nothing in this agenda assumes prior concurrency experience; it "
        "assumes curiosity about why concurrent code is harder than it "
        "looks, and gives that curiosity a precise vocabulary to work with.")


def build_closing(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.BG_TITLE_DARK)
    bp.add_text(slide, Inches(1.0), Inches(1.6), Inches(11.3), Inches(1.0),
                "The model, not the API.", size=F.TITLE_LARGE, color=C.TEXT_WHITE,
                bold=True, font=FontFam.HEADER)
    bp.add_multi_text(slide, Inches(1.0), Inches(2.8), Inches(11.3), Inches(2.8), [
        dict(text="Concurrency is structure; parallelism is execution. Keep the two axes apart "
                  "and most sizing mistakes stop happening.", size=Pt(18), color=C.TEXT_LIGHT,
             font=FontFam.BODY),
        dict(text="Shared state is hard because the hardware and the compiler are both allowed "
                  "to reorder what you did not explicitly order.", size=Pt(18),
             color=C.TEXT_LIGHT, font=FontFam.BODY, space_before=Pt(10)),
        dict(text="Trust a number only after it survives warmup, coordinated omission, and a "
                  "percentile instead of a mean — and remember the container is watching too.",
             size=Pt(18), color=C.TEXT_LIGHT, font=FontFam.BODY, space_before=Pt(10)),
        dict(text="Next: the 101 deck builds the C++ toolkit; the 201 deck tunes it for "
                  "microseconds.", size=Pt(18), color=C.ACCENT_CYAN,
             font=FontFam.BODY, space_before=Pt(10)),
    ])
    bp.add_text(slide, Inches(1.0), Inches(6.0), Inches(11.3), Inches(0.5),
                "Read the book  •  run the examples  •  the source is on the repo",
                size=F.BODY, color=C.TEXT_MUTED, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(6.6), Inches(11.3), Inches(0.4),
                REPO_URL, size=F.CAPTION, color=C.TEXT_MUTED, font=FontFam.CODE)
    bp.set_notes(slide,
        "To close: this deck's entire job was to replace a vague feeling "
        "that concurrency is hard with a precise model of why it is hard. "
        "Structure and execution are independent axes, the machine "
        "underneath a thread has real, measurable costs that a correct "
        "program can still pay unnecessarily, shared state needs both "
        "atomicity and ordering, and a performance number is only "
        "trustworthy once it has survived warmup, coordinated omission, and "
        "a percentile report rather than a mean.\n\n"
        "None of this required much C++ code, on purpose. The 101 deck "
        "picks up immediately after this one and builds the concrete "
        "toolkit, threads, coroutines, futures, lock-free structures, "
        "parallel algorithms, and the 201 deck tunes that toolkit for "
        "low-latency, high-throughput systems where the concepts in this "
        "deck stop being background and become the daily unit of reasoning. "
        "Thank you.")


def main():
    sys.path.insert(0, str(TOOLS))
    from sections_conc_100 import SECTIONS

    total = 2  # title + agenda
    for s in SECTIONS:
        total += 1 + len(s["slides"])
    total += 1  # closing
    print(f"Total slides planned: {total}")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    page = 1
    print(f"  [{page:3d}/{total}] Title"); build_title(prs, total); page += 1
    print(f"  [{page:3d}/{total}] Agenda"); build_agenda(prs, total, SECTIONS); page += 1

    for section in SECTIONS:
        print(f"  [{page:3d}/{total}] §{section['num']} divider — {section['title']}")
        bp.build_section_divider(prs, section["num"], section["title"],
                                 section["tagline"], page, total,
                                 section["divider_notes"])
        page += 1
        for slide_data in section["slides"]:
            kind = slide_data["kind"]
            t = slide_data.get("title") or slide_data.get("demo_name", "(slide)")
            print(f"  [{page:3d}/{total}] §{section['num']} {kind} — {t[:48]}")
            bp.render_slide(prs, section, slide_data, page, total)
            page += 1

    print(f"  [{page:3d}/{total}] Closing"); build_closing(prs, total); page += 1

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PATH))
    print(f"\nSaved: {OUT_PATH}")
    print(f"Slide count: {page - 1}")
    bp.write_deck_notes(NOTES_PATH, "100", [
        "## cpp-concurrency-100-mental-model.pptx (the 100 deck)",
        "",
        f"- Total slides: {page - 1}",
        f"- Sections rendered: {len(SECTIONS)}",
        "- Generator: tools/build-conc-100-pptx.py",
        "- Content source: tools/sections_conc_100.py",
        "- Shared builders: tools/build-pptx.py (loaded via importlib)",
        "- Audience: engineers new to concurrency, tech leads scoping a system",
        "- Diagrams embedded: 01-execution-models, 02-cache-hierarchy-mesi, "
        "02-context-switch-anatomy, 03-clock-taxonomy, 03-deadline-propagation, "
        "06-lock-ordering-deadlock, 08-memory-ordering, 18-coordinated-omission, "
        "22-communication-families, 27-tsan-report-anatomy, 32-quota-vs-cpuset",
    ])
    print(f"Wrote: {NOTES_PATH}")


if __name__ == "__main__":
    main()
