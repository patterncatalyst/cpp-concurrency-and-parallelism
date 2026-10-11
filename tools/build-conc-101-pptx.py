#!/usr/bin/env python3
"""
build-conc-101-pptx.py -- render the C++ CONCURRENCY 101 deck
("The Modern C++ Concurrency Toolkit").

Thin entry script mirroring the statelessness-101 precedent. It reuses every
slide builder, theme, and helper from build-pptx.py -- loaded via importlib
because that file's name has a hyphen -- and supplies its own title / agenda
/ closing slides plus the twelve-chapter SECTIONS from sections_conc_101.py.

This is the code-heavy on-ramp deck for the C++ concurrency toolkit:
threads, locks, coordination primitives, atomics, lock-free structures,
futures, coroutines, async runtimes, fibers, parallel algorithms, and thread
pools. It does not cover low-latency/HFT-specific material -- that is the
201 deck's job, pointed to from the closing slide.

Output: presentation/cpp-concurrency-101-cpp-toolkit.pptx
"""
import importlib.util
import sys
from pathlib import Path

from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
TOOLS = ROOT / "tools"


def load_build_pptx():
    """Load build-pptx.py (hyphenated filename) as a module."""
    spec = importlib.util.spec_from_file_location(
        "build_pptx", TOOLS / "build-pptx.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


bp = load_build_pptx()
bp.FOOTER_LABEL = "C++ Concurrency 101"
C, F, FontFam = bp.C, bp.F, bp.FontFam
SLIDE_W, SLIDE_H = bp.SLIDE_W, bp.SLIDE_H

OUT_PATH = ROOT / "presentation" / "cpp-concurrency-101-cpp-toolkit.pptx"
NOTES_PATH = ROOT / "presentation" / "build-notes-conc.md"
REPO_URL = "github.com/patterncatalyst/cpp-concurrency-and-parallelism"


def build_title(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.BG_TITLE_DARK)

    dot_y = Inches(1.5)
    for i, color in enumerate([C.ACCENT_BLUE, C.ACCENT_GREEN, C.ACCENT_CYAN]):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(1.0 + i * 0.35), dot_y, Inches(0.22), Inches(0.22))
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()

    bp.add_text(slide, Inches(1.0), Inches(2.0), Inches(11.3), Inches(0.5),
                "C++ CONCURRENCY 101", size=Pt(20), color=C.ACCENT_BLUE,
                bold=True, font=FontFam.HEADER)
    bp.add_text(slide, Inches(1.0), Inches(2.5), Inches(11.3), Inches(1.5),
                "The Modern C++ Concurrency Toolkit",
                size=F.TITLE_HUGE, color=C.TEXT_WHITE, bold=True, font=FontFam.HEADER)
    bp.add_text(slide, Inches(1.0), Inches(4.5), Inches(11.3), Inches(0.5),
                "jthread + stop_token  •  mutexes, locks, lock ordering  •  "
                "atomics & the memory model  •  lock-free structures",
                size=F.BODY, color=C.TEXT_LIGHT, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(5.0), Inches(11.3), Inches(0.5),
                "futures & coroutines  •  async runtimes  •  fibers  •  "
                "parallel algorithms  •  thread pools",
                size=F.BODY, color=C.TEXT_LIGHT, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(5.7), Inches(11.3), Inches(0.4),
                "Twelve chapters of real, compiling C++20/23/26 — what "
                "ships today and what's still experimental",
                size=F.BODY_SMALL, color=C.TEXT_MUTED, italic=True, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.4),
                REPO_URL, size=F.CAPTION, color=C.TEXT_MUTED, font=FontFam.CODE)

    bp.set_notes(slide,
        "Welcome to the modern C++ concurrency toolkit. This deck is "
        "code-heavy and broad: twelve chapters' worth of the "
        "standard library's concurrency surface, from std::jthread through "
        "thread pools, with real, compiling C++ on every slide.\n\n"
        "The organizing idea is the same one the book uses: show what the "
        "standard actually guarantees, show the gotcha that trips people up "
        "in practice, and name the standards delta precisely -- what's "
        "C++20, what C++23 only refines, and what C++26 proposes but no "
        "compiler ships yet. Treat every 'proposed' and 'experimental' "
        "label in this deck as accurate, not decoration.\n\n"
        "This is the toolkit tour, not the low-latency playbook. The 201 "
        "deck takes this same vocabulary and applies it under a real "
        "latency budget -- mechanical sympathy, lock-free messaging, the "
        "disruptor pattern. That's next; today is what the pieces are and "
        "how they actually behave.")


def build_agenda(prs, total, sections):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.TEXT_WHITE)
    bp.add_header_bar(slide, "Overview", 2, total, "Agenda — the twelve chapters")
    bp.add_footer(slide, "Agenda", 2, total)

    palette = [C.ACCENT_RED, C.ACCENT_ORANGE, C.ACCENT_BLUE, C.ACCENT_GREEN,
               C.ACCENT_CYAN, C.ACCENT_PURPLE]
    rows_per_col = 6
    col_width = Inches(6.0)
    row_height = Inches(0.62)
    grid_left = Inches(0.5)
    grid_top = Inches(1.1)

    for i, sec in enumerate(sections):
        color = palette[i % len(palette)]
        col = i // rows_per_col
        row = i % rows_per_col
        x = grid_left + col * (col_width + Inches(0.3))
        y = grid_top + row * row_height

        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, Inches(0.5), Inches(0.5))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
        bp.add_text(slide, x, y, Inches(0.5), Inches(0.5), str(sec["num"]),
                    size=Pt(16), color=C.TEXT_WHITE, bold=True,
                    align=PP_ALIGN.CENTER, anchor=bp.MSO_ANCHOR.MIDDLE,
                    font=FontFam.HEADER)
        bp.add_text(slide, x + Inches(0.65), y, Inches(5.2), Inches(0.5),
                    sec["title"], size=Pt(15), color=C.TEXT_DARK,
                    anchor=bp.MSO_ANCHOR.MIDDLE, font=FontFam.BODY)

    bp.set_notes(slide,
        "Twelve sections, one per source chapter, in the book's own order. "
        "The first six build the shared-state foundation: threads and "
        "cooperative cancellation, the futex underneath the mutex, the "
        "mutex family and lock ordering, coordination primitives, atomics "
        "and the memory model, and lock-free structures with the ABA "
        "problem.\n\n"
        "The next three move to composition: futures and promises, "
        "coroutines in depth, and async runtimes built on epoll and "
        "Boost.Asio. Fibers close out the suspension story as the stackful "
        "alternative to a stackless coroutine.\n\n"
        "The last two turn from suspension to raw throughput: parallel "
        "algorithms and execution policies, then thread pools and work "
        "scheduling as the longer-lived structure that holds everything "
        "else. Every section ends with a standards-delta note -- what's "
        "C++20, what C++23 changed, and what C++26 only proposes.")


def build_closing(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.BG_TITLE_DARK)
    bp.add_text(slide, Inches(1.0), Inches(1.6), Inches(11.3), Inches(1.0),
                "The toolkit is only half the job.", size=F.TITLE_LARGE, color=C.TEXT_WHITE,
                bold=True, font=FontFam.HEADER)
    bp.add_multi_text(slide, Inches(1.0), Inches(2.8), Inches(11.3), Inches(2.8), [
        dict(text="Every primitive in this deck compiles correctly long before it "
                  "performs correctly under a real latency budget.", size=Pt(18),
             color=C.TEXT_LIGHT, font=FontFam.BODY),
        dict(text="jthread, the memory model, lock-free reclamation, and parallel "
                  "algorithms are the vocabulary -- mechanical sympathy and the "
                  "disruptor pattern are what turn that vocabulary into microseconds.",
             size=Pt(18), color=C.TEXT_LIGHT, font=FontFam.BODY, space_before=Pt(10)),
        dict(text="Next: the 201 deck -- low-latency and HFT patterns built on exactly "
                  "the primitives covered today.", size=Pt(18),
             color=C.ACCENT_CYAN, font=FontFam.BODY, space_before=Pt(10)),
    ])
    bp.add_text(slide, Inches(1.0), Inches(5.9), Inches(11.3), Inches(0.5),
                "Read the book  •  run the chapter examples  •  the source is on the repo",
                size=F.BODY, color=C.TEXT_MUTED, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.4),
                REPO_URL, size=F.CAPTION, color=C.TEXT_MUTED, font=FontFam.CODE)
    bp.set_notes(slide,
        "That's the toolkit: twelve chapters of standard-library concurrency, "
        "with the standards delta named precisely at every stop so nothing "
        "proposed gets mistaken for something shipping.\n\n"
        "The throughline worth keeping: correctness and performance are "
        "separate questions here. A jthread joins correctly, a release/"
        "acquire pair hands off correctly, a work-stealing pool balances "
        "load correctly -- none of that says anything yet about whether the "
        "result meets a microsecond budget.\n\n"
        "The 201 deck is where this toolkit meets that budget: mechanical "
        "sympathy, the low-latency fast path, lock-free messaging, and the "
        "disruptor pattern, built on exactly the primitives covered today. "
        "Thanks for sitting through the tour.")


def main():
    sys.path.insert(0, str(TOOLS))
    from sections_conc_101 import SECTIONS

    total = 2  # title + agenda
    for s in SECTIONS:
        total += 1 + len(s["slides"])
    total += 1  # closing
    print(f"Total slides planned: {total}")

    prs = bp.Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    page = 1
    print(f"  [{page:3d}/{total}] Title"); build_title(prs, total); page += 1
    print(f"  [{page:3d}/{total}] Agenda"); build_agenda(prs, total, SECTIONS); page += 1

    for section in SECTIONS:
        print(f"  [{page:3d}/{total}] §{section['num']} divider — {section['title']}")
        bp.build_section_divider(prs, section["num"], section["title"],
                                 section["tagline"], page, total,
                                 section["divider_notes"])
        page += 1
        for slide_data in section["slides"]:
            kind = slide_data["kind"]
            t = slide_data.get("title") or slide_data.get("demo_name", "(demo)")
            print(f"  [{page:3d}/{total}] §{section['num']} {kind} — {t[:48]}")
            bp.render_slide(prs, section, slide_data, page, total)
            page += 1

    print(f"  [{page:3d}/{total}] Closing"); build_closing(prs, total); page += 1

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PATH))
    print(f"\nSaved: {OUT_PATH}")
    print(f"Slide count: {page - 1}")
    bp.write_deck_notes(NOTES_PATH, "101", [
        "## cpp-concurrency-101-cpp-toolkit.pptx (the 101 deck)",
        "",
        f"- Total slides: {page - 1}",
        f"- Sections rendered: {len(SECTIONS)}",
        "- Generator: tools/build-conc-101-pptx.py",
        "- Content source: tools/sections_conc_101.py",
        "- Shared builders: tools/build-pptx.py (loaded via importlib)",
    ])
    print(f"Wrote: {NOTES_PATH}")


if __name__ == "__main__":
    main()
