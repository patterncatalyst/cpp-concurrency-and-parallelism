#!/usr/bin/env python3
"""
build-conc-201-pptx.py — render the LOW-LATENCY & HIGH-THROUGHPUT 201 deck.

Companion to build-pptx.py (the shared engine). Reuses every slide builder,
the theme, and the helpers from build-pptx.py — loaded via importlib because
that file's name has a hyphen — and supplies its own title / agenda / closing
slides plus the SECTIONS content from sections_conc_201.py.

Diagrams: tools/convert-diagrams.sh converts diagrams/*.svg into
/tmp/diagrams-png/<name>.jpg before this runs.

Output: presentation/cpp-concurrency-201-low-latency.pptx
"""
import importlib.util
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent.parent
TOOLS = ROOT / "tools"


def load_build_pptx():
    """Load build-pptx.py (hyphenated filename) as a module."""
    spec = importlib.util.spec_from_file_location(
        "build_pptx", TOOLS / "build-pptx.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


bp = load_build_pptx()
bp.FOOTER_LABEL = "C++ Concurrency 201"
C, F, FontFam = bp.C, bp.F, bp.FontFam
SLIDE_W, SLIDE_H = bp.SLIDE_W, bp.SLIDE_H

OUT_PATH = ROOT / "presentation" / "cpp-concurrency-201-low-latency.pptx"
NOTES_PATH = ROOT / "presentation" / "build-notes-conc.md"
REPO_URL = "github.com/patterncatalyst/cpp-concurrency-and-parallelism"


def build_title(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.BG_TITLE_DARK)

    dot_y = Inches(1.5)
    for i, color in enumerate([C.ACCENT_CYAN, C.ACCENT_GREEN, C.ACCENT_ORANGE]):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(1.0 + i * 0.35), dot_y, Inches(0.22), Inches(0.22))
        dot.fill.solid()
        dot.fill.fore_color.rgb = color
        dot.line.fill.background()

    bp.add_text(slide, Inches(1.0), Inches(2.0), Inches(11.3), Inches(0.5),
                "LOW-LATENCY & HIGH-THROUGHPUT C++ — THE 201 DEEP DIVE",
                size=Pt(20), color=C.ACCENT_CYAN,
                bold=True, font=FontFam.HEADER)
    bp.add_text(slide, Inches(1.0), Inches(2.5), Inches(11.3), Inches(1.5),
                "Mechanical Sympathy, Lock-Free Messaging,\nand the Pipeline That Ties Them Together",
                size=F.TITLE_HUGE, color=C.TEXT_WHITE, bold=True, font=FontFam.HEADER)
    bp.add_text(slide, Inches(1.0), Inches(4.5), Inches(11.3), Inches(0.5),
                "NUMA and SIMD  •  benchmarking without deception  •  false sharing  "
                "•  the Disruptor  •  observability  •  the full pipeline",
                size=F.BODY, color=C.TEXT_LIGHT, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(5.2), Inches(11.3), Inches(0.4),
                "Companion deck to cpp-concurrency-and-parallelism  —  Docs 16–33",
                size=F.BODY_SMALL, color=C.TEXT_MUTED, italic=True, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(5.9), Inches(11.3), Inches(0.4),
                "Every pattern here has a runnable demo under examples/concurrency/",
                size=F.BODY_SMALL, color=C.ACCENT_CYAN, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(6.5), Inches(11.3), Inches(0.4),
                REPO_URL, size=F.CAPTION, color=C.TEXT_MUTED, font=FontFam.CODE)

    bp.set_notes(slide,
        "This is the 201 deck — the low-latency and high-throughput deep dive for "
        "senior and staff engineers building matching engines, feed handlers, and "
        "other systems where single-digit microseconds are the budget, not an "
        "aspiration. Where an introductory deck explains what a thread is, this one "
        "assumes that and goes straight at NUMA placement, false sharing, the "
        "Disruptor pattern, and a full pipeline that composes all of it under "
        "load.\n\n"
        "The spine is fourteen sections built from chapters 16 through 33 of the "
        "book: locality and vectorization, the benchmarking discipline that makes "
        "every later performance claim trustworthy, the fast path and mechanical "
        "sympathy, lock-free messaging, communication models and the actor model, "
        "concurrent REST and gRPC under load, error handling across thread "
        "boundaries, debugging and sanitizers, observability, coding practices and "
        "architecture patterns, running correctly in a container, and a full end-to-end example "
        "order-book pipeline that ties the whole arc together.\n\n"
        "Every pattern in this deck has a runnable companion under examples/"
        "concurrency/, and we will cue each demo at the point its pattern is "
        "introduced — eight of them across the deck, from NUMA placement in a "
        "disposable KVM lab through the full ingest-match-publish pipeline measured "
        "under a deliberate overload burst.")


def build_agenda(prs, total, sections):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.TEXT_WHITE)
    bp.add_header_bar(slide, "Overview", 2, total, "Agenda — fourteen sections")
    bp.add_footer(slide, "Agenda", 2, total)

    palette = [C.ACCENT_RED, C.ACCENT_ORANGE, C.ACCENT_BLUE, C.ACCENT_GREEN,
               C.ACCENT_CYAN, C.ACCENT_PURPLE]
    rows_per_col = 7
    col_width = Inches(6.0)
    row_height = Inches(0.62)
    grid_left = Inches(0.5)
    grid_top = Inches(1.1)

    for i, sec in enumerate(sections):
        color = palette[i % len(palette)]
        col = i // rows_per_col
        row = i % rows_per_col
        x = grid_left + col * (col_width + Inches(0.3))
        y = grid_top + row * row_height

        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, Inches(0.5), Inches(0.5))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
        bp.add_text(slide, x, y, Inches(0.5), Inches(0.5), str(sec["num"]),
                    size=Pt(16), color=C.TEXT_WHITE, bold=True,
                    align=PP_ALIGN.CENTER, anchor=bp.MSO_ANCHOR.MIDDLE,
                    font=FontFam.HEADER)
        bp.add_text(slide, x + Inches(0.65), y, Inches(5.2), Inches(0.5),
                    sec["title"], size=Pt(14), color=C.TEXT_DARK,
                    anchor=bp.MSO_ANCHOR.MIDDLE, font=FontFam.BODY)

    bp.set_notes(slide,
        "Fourteen sections, two movements. Sections one through six build the "
        "low-latency foundation — locality, vectorization, benchmarking "
        "methodology, the fast path, mechanical sympathy, and lock-free "
        "messaging. Sections seven through thirteen widen into throughput and "
        "correctness at scale — communication models and actors, concurrent REST "
        "and gRPC, error handling across boundaries, debugging and sanitizers, "
        "observability, coding and architecture patterns, and running correctly "
        "under a container's real resource limits.\n\n"
        "Section fourteen closes with the full example: a single ingest-match-publish "
        "pipeline that composes a thread pool, a pinned single-writer thread, "
        "lock-free rings, and full tracing into one measured system, proving the "
        "earlier sections' individually-true claims hold up stacked together. "
        "Eight sections pause on a demo cue where the pattern has a runnable "
        "companion.")


def build_closing(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bp.set_slide_bg(slide, C.BG_TITLE_DARK)
    bp.add_text(slide, Inches(1.0), Inches(1.6), Inches(11.3), Inches(1.0),
                "Every claim here is runnable.", size=F.TITLE_LARGE, color=C.TEXT_WHITE,
                bold=True, font=FontFam.HEADER)
    bp.add_multi_text(slide, Inches(1.0), Inches(2.8), Inches(11.3), Inches(2.6), [
        dict(text="Locality, vectorization, and a trustworthy benchmark are the "
                  "foundation everything else in this deck stands on.", size=Pt(18),
             color=C.TEXT_LIGHT, font=FontFam.BODY),
        dict(text="Mechanical sympathy and the Disruptor turn that foundation into "
                  "a fast path; communication models and architecture patterns turn "
                  "the fast path into a whole system.", size=Pt(18),
             color=C.TEXT_LIGHT, font=FontFam.BODY, space_before=Pt(10)),
        dict(text="The full pipeline composes all of it, measured under load — "
                  "and eight demos across this deck are where a reader's own hardware "
                  "produces the real numbers.", size=Pt(18), color=C.ACCENT_CYAN,
             font=FontFam.BODY, space_before=Pt(10)),
    ])
    bp.add_text(slide, Inches(1.0), Inches(5.6), Inches(11.3), Inches(0.5),
                "Read the book  •  run the demos under examples/concurrency/  •  the source is on the repo",
                size=F.BODY, color=C.TEXT_MUTED, font=FontFam.BODY)
    bp.add_text(slide, Inches(1.0), Inches(6.4), Inches(11.3), Inches(0.4),
                REPO_URL, size=F.CAPTION, color=C.TEXT_MUTED, font=FontFam.CODE)
    bp.set_notes(slide,
        "To close: every pattern in this deck, from first-touch NUMA placement to "
        "the full example's pinned single-writer match thread, is backed by a runnable "
        "example under examples/concurrency/, and none of the performance claims in "
        "this deck are presented as the book's own measurement — the hardware a "
        "reader runs these on produces the real figures.\n\n"
        "The arc holds together end to end: a trustworthy benchmark makes every "
        "other number in the book meaningful, mechanical sympathy and lock-free "
        "messaging turn that discipline into a fast path, and the full pipeline "
        "proves the individually-true claims from every earlier section still hold "
        "once they are stacked into one system under real load. Clone the repo, run "
        "the demos, and measure your own hardware. Thank you.")


def main():
    sys.path.insert(0, str(TOOLS))
    from sections_conc_201 import SECTIONS

    total = 2  # title + agenda
    for s in SECTIONS:
        total += 1 + len(s["slides"])
    total += 1  # closing
    print(f"Total slides planned: {total}")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    page = 1
    print(f"  [{page:3d}/{total}] Title"); build_title(prs, total); page += 1
    print(f"  [{page:3d}/{total}] Agenda"); build_agenda(prs, total, SECTIONS); page += 1

    for section in SECTIONS:
        print(f"  [{page:3d}/{total}] §{section['num']} divider — {section['title']}")
        bp.build_section_divider(prs, section["num"], section["title"],
                                 section["tagline"], page, total,
                                 section["divider_notes"])
        page += 1
        for slide_data in section["slides"]:
            kind = slide_data["kind"]
            t = slide_data.get("title") or slide_data.get("demo_name", "(demo)")
            print(f"  [{page:3d}/{total}] §{section['num']} {kind} — {t[:48]}")
            bp.render_slide(prs, section, slide_data, page, total)
            page += 1

    print(f"  [{page:3d}/{total}] Closing"); build_closing(prs, total); page += 1

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PATH))
    print(f"\nSaved: {OUT_PATH}")
    print(f"Slide count: {page - 1}")
    bp.write_deck_notes(NOTES_PATH, "201", [
        "## cpp-concurrency-201-low-latency.pptx (the 201 deck)",
        "",
        f"- Total slides: {page - 1}",
        f"- Sections rendered: {len(SECTIONS)}",
        "- Generator: tools/build-conc-201-pptx.py",
        "- Content source: tools/sections_conc_201.py",
        "- Shared builders: tools/build-pptx.py (loaded via importlib)",
    ])
    print(f"Wrote: {NOTES_PATH}")


if __name__ == "__main__":
    main()
